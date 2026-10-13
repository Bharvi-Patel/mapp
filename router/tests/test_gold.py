import os, sys
HERE = os.path.dirname(__file__)
sys.path.insert(0, os.path.join(HERE, '..'))
import gold

CFG = {'sku_prefix_rules': {'100-': 'store-displays'},
       'tie_prefer': [{'prefer': 'mattresses-by-size-', 'over': 'mattresses-and-bases-by-type-'}]}
STORE = {'rugs-area-rugs': True, 'rugs-washable': True, 'nursery-cribs': False,
         'mattresses-by-size-queen': True, 'mattresses-and-bases-by-type-hybrid': True}
def L(slug, u=''): return {'slug': slug, 'unsure': u}
def d(sku, a, b, ruling=None): return gold.decide(sku, a, b, CFG, STORE, ruling)

def test_sku_prefix_rule_wins():
    assert d('100-12', L('__none__'), L('__none__'))[:3] == ('rule_handled', 'store-displays', 'sku_prefix:100-')

def test_agree_gold_and_flagged():
    assert d('A1', L('rugs-area-rugs'), L('rugs-area-rugs'))[:2] == ('gold', 'rugs-area-rugs')
    assert d('A1', L('rugs-area-rugs', 'Y'), L('rugs-area-rugs'))[0] == 'gold_flagged'

def test_none_is_a_gold_answer():
    assert d('A1', L('__none__'), L('__none__'))[:2] == ('gold', '__none__')

def test_category_not_in_store_is_held():
    assert d('A1', L('nursery-cribs'), L('nursery-cribs'))[0] == 'hold_not_in_store'

def test_split_goes_to_review_unless_ruled():
    assert d('A1', L('rugs-area-rugs'), L('rugs-washable'))[0] == 'review'
    st, slug, rule, _ = d('A1', L('rugs-area-rugs'), L('rugs-washable'), {'slug': 'rugs-washable', 'reason': 'x'})
    assert (st, slug, rule) == ('gold_ruled', 'rugs-washable', 'ruling')

def test_tie_prefer_either_order():
    for a, b in (('mattresses-by-size-queen', 'mattresses-and-bases-by-type-hybrid'),
                 ('mattresses-and-bases-by-type-hybrid', 'mattresses-by-size-queen')):
        assert d('A1', L(a), L(b))[:2] == ('gold_ruled', 'mattresses-by-size-queen')

def test_wilson_and_slug_normalising():
    lo, hi = gold.wilson(57, 60)
    assert 0.85 < lo < 0.87 and 0.98 < hi < 1.0
    assert gold.norm_slug('Rugs   |   Rugs > Washable   [rugs-washable]') == 'rugs-washable'
    assert gold.norm_slug('NONE FITS') == '__none__' and gold.norm_slug('') == ''

def test_spot_xlsx_roundtrip(tmp_path):
    import openpyxl
    t = tmp_path / 'sheet.xlsx'
    wb = openpyxl.Workbook(); wb.active.title = 'Read me'; wb['Read me']['A1'] = 'x'
    cat = wb.create_sheet('Categories'); cat['A1'] = 'h'; cat['A2'] = 'NONE FITS'; cat['A3'] = 'Rug   |   Rugs   [rugs-washable]'
    ws = wb.create_sheet('Label'); ws['A3'] = 'Row ID'; ws['H3'] = 'Your category'; ws['I3'] = 'Unsure (Y)'
    for k in range(4):
        ws.cell(4 + k, 1).value = f'T000{k + 1}'; ws.cell(4 + k, 2).value = f'SKU{k}'; ws.cell(4 + k, 3).value = f'name{k}'
    dv = gold.DataValidation(type='list', formula1='=Categories!$A$2:$A$3'); dv.add('H4:H7'); ws.add_data_validation(dv)
    wb.save(t)
    out = tmp_path / 'spot.xlsx'
    gold.write_spot_xlsx(str(t), ['T0003', 'T0001'], str(out))
    w2 = openpyxl.load_workbook(out); l2 = w2['Label']
    assert [l2.cell(4, 1).value, l2.cell(5, 1).value, l2.cell(6, 1).value] == ['T0003', 'T0001', None]
    assert l2['D3'].value == 'Your category' and l2['C4'].value == 'name2'
    assert any(str(d.sqref) == 'D4:D5' for d in l2.data_validations.dataValidation)
    l2['D4'] = 'Rug   |   Rugs   [rugs-washable]'; w2.save(out)
    assert gold.read_human(str(out)) == [{'row_id': 'T0003', 'human_slug': 'Rug   |   Rugs   [rugs-washable]'},
                                         {'row_id': 'T0001', 'human_slug': ''}]

def test_classify_answers():
    leaf = {'rugs-area-rugs', 'rugs-washable', 'sofas-and-seating-sofas'}
    assert gold.classify('rugs-area-rugs', 'rugs-area-rugs', leaf) == 'match'
    assert gold.classify('rugs-area-rugs', 'rugs-and-decor-rugs', leaf) == 'nonleaf'     # a parent, not a leaf
    assert gold.classify('rugs-area-rugs', '__none__', leaf) == 'human_none'
    assert gold.classify('__none__', 'sofas-and-seating-sofas', leaf) == 'key_none'
    assert gold.classify('rugs-area-rugs', 'rugs-washable', leaf) == 'other'

def test_exclude_ruling_leaves_row_out():
    st, slug, rule, note = d('A1', L('rugs-area-rugs'), L('rugs-area-rugs'), {'slug': '__exclude__', 'reason': 'contested'})
    assert (st, slug, rule) == ('excluded', '', 'ruling')